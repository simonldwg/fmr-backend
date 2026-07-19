"""
Django management command: Exports statistics about the music library
as an Excel file (.xlsx).

Usage:
    python manage.py export_stats
    python manage.py export_stats --output stats.xlsx --bins-arousal 15 --bins-bpm 20
"""

from collections import Counter

from django.core.management.base import BaseCommand
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from apps.library.models import Song


class Command(BaseCommand):
    help = "Exports statistics (genres, arousal, BPM, top artists) as an Excel file."

    def add_arguments(self, parser):
        parser.add_argument(
            "--output",
            default="song_stats.xlsx",
            help="Path of the output file (default: song_stats.xlsx)",
        )
        parser.add_argument(
            "--bins-arousal",
            type=int,
            default=10,
            help="Number of bins for the arousal distribution (default: 10)",
        )
        parser.add_argument(
            "--bins-bpm",
            type=int,
            default=10,
            help="Number of bins for the BPM distribution (default: 10)",
        )

    def handle(self, *args, **options):
        output_path = options["output"]
        bins_arousal = options["bins_arousal"]
        bins_bpm = options["bins_bpm"]

        songs = Song.objects.select_related("features", "genres")
        total = songs.count()
        if total == 0:
            self.stdout.write(self.style.WARNING("No songs in the database."))
            return

        genre_counter = Counter()
        artist_counter = Counter()
        arousal_values = []
        bpm_values = []

        for song in songs.iterator():
            artist_counter[song.artist] += 1

            top3 = song.genres.top3_genres if song.genres else None
            if top3:
                genre_counter.update(top3.keys())

            if song.features:
                if song.features.arousal is not None:
                    arousal_values.append(song.features.arousal)
                if song.features.bpm is not None:
                    bpm_values.append(song.features.bpm)

        wb = Workbook()
        wb.remove(wb.worksheets[0])  # remove default sheet

        self._write_count_sheet(
            wb,
            title="Genres",
            header=("Genre", "Count"),
            rows=genre_counter.most_common(),
            total=total,
            chart_title="Genre frequency (top 3 genres per song)",
        )

        self._write_histogram_sheet(
            wb,
            title="Arousal Distribution",
            values=arousal_values,
            n_bins=bins_arousal,
            value_label="Arousal",
            decimals=2,
        )

        self._write_histogram_sheet(
            wb,
            title="BPM Distribution",
            values=bpm_values,
            n_bins=bins_bpm,
            value_label="BPM",
            decimals=1,
        )

        self._write_count_sheet(
            wb,
            title="Top Artists",
            header=("Artist", "Song count"),
            rows=artist_counter.most_common(10),
            total=total,
            chart_title="Ten most frequent artists",
        )

        wb.save(output_path)
        self.stdout.write(
            self.style.SUCCESS(f"Analyzed {total} songs, file saved: {output_path}")
        )

    def _write_count_sheet(self, wb, title, header, rows, total, chart_title):
        """Sheet with category/count/share + bar chart."""
        ws = wb.create_sheet(title)
        ws.append([*header, "Share"])
        for cell in ws[1]:
            cell.font = Font(bold=True)

        for name, count in rows:
            ws.append([name, count, count / total])

        for row in range(2, ws.max_row + 1):
            ws.cell(row=row, column=3).number_format = "0.0%"

        self._autosize(ws)
        self._add_bar_chart(ws, chart_title, n_rows=len(rows), anchor="E2")

    def _write_histogram_sheet(self, wb, title, values, n_bins, value_label, decimals):
        """Sheet with histogram bins + summary statistics + bar chart."""
        ws = wb.create_sheet(title)

        if not values:
            ws.append(["No data available."])
            return

        vmin, vmax = min(values), max(values)
        if vmin == vmax:
            bins = [(vmin, vmax, len(values))]
        else:
            width = (vmax - vmin) / n_bins
            counts = [0] * n_bins
            for v in values:
                idx = min(int((v - vmin) / width), n_bins - 1)
                counts[idx] += 1
            bins = [
                (vmin + i * width, vmin + (i + 1) * width, counts[i])
                for i in range(n_bins)
            ]

        ws.append(["Bin", "Count", "Share"])
        for cell in ws[1]:
            cell.font = Font(bold=True)

        fmt = f"{{:.{decimals}f}}"
        n = len(values)
        for lo, hi, count in bins:
            label = f"{fmt.format(lo)} – {fmt.format(hi)}"
            ws.append([label, count, count / n])

        for row in range(2, ws.max_row + 1):
            ws.cell(row=row, column=3).number_format = "0.0%"

        mean = sum(values) / n
        sorted_vals = sorted(values)
        mid = n // 2
        median = sorted_vals[mid] if n % 2 else (sorted_vals[mid - 1] + sorted_vals[mid]) / 2
        stats = [
            ("Count", n),
            ("Minimum", vmin),
            ("Maximum", vmax),
            ("Mean", mean),
            ("Median", median),
        ]
        ws.cell(row=1, column=5, value=f"{value_label} summary").font = Font(bold=True)
        for i, (label, value) in enumerate(stats, start=2):
            ws.cell(row=i, column=5, value=label)
            ws.cell(row=i, column=6, value=round(value, 4))

        self._autosize(ws)
        self._add_bar_chart(
            ws, f"{value_label} distribution", n_rows=len(bins), anchor="E9"
        )

    @staticmethod
    def _add_bar_chart(ws, title, n_rows, anchor):
        chart = BarChart()
        chart.title = title
        chart.legend = None
        chart.y_axis.title = "Count"
        data = Reference(ws, min_col=2, min_row=1, max_row=n_rows + 1)
        cats = Reference(ws, min_col=1, min_row=2, max_row=n_rows + 1)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.width = 20
        chart.height = 10
        ws.add_chart(chart, anchor)

    @staticmethod
    def _autosize(ws):
        for col_cells in ws.iter_cols(min_col=1, max_col=3):
            length = max(
                (len(str(c.value)) for c in col_cells if c.value is not None), default=0
            )
            ws.column_dimensions[get_column_letter(col_cells[0].column)].width = length + 2
